#!/usr/bin/env python3
"""
Extract text content from PowerPoint files and format for ISA context addition.
"""

import sys
import os
from pptx import Presentation
from pathlib import Path

def extract_text_from_pptx(pptx_path):
    """Extract all text content from a PowerPoint file."""
    try:
        prs = Presentation(pptx_path)
        content = []
        
        # Add title
        filename = Path(pptx_path).stem
        content.append(f"# {filename}")
        content.append("")
        
        for slide_num, slide in enumerate(prs.slides, 1):
            content.append(f"## Slide {slide_num}")
            content.append("")
            
            # Extract text from all shapes in the slide
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
            
            if slide_text:
                content.extend(slide_text)
            else:
                content.append("(No text content)")
            
            content.append("")
        
        return "\n".join(content)
    
    except Exception as e:
        return f"Error processing {pptx_path}: {str(e)}"

def main():
    if len(sys.argv) < 2:
        print("Usage: python extract_pptx_for_isa.py <pptx_file> [output_file]")
        sys.exit(1)
    
    pptx_path = sys.argv[1]
    
    if not os.path.exists(pptx_path):
        print(f"Error: File '{pptx_path}' not found")
        sys.exit(1)
    
    # Extract content
    content = extract_text_from_pptx(pptx_path)
    
    # Determine output file
    if len(sys.argv) > 2:
        output_path = sys.argv[2]
    else:
        # Create output filename based on input
        input_path = Path(pptx_path)
        output_path = input_path.parent / f"{input_path.stem}_extracted.md"
    
    # Write to file
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(content)
    
    print(f"Extracted content from '{pptx_path}' to '{output_path}'")
    print(f"You can now add to ISA with: isa ctx-add '{output_path}'")

if __name__ == "__main__":
    main()